"""
Catalog ingestion — run this manually whenever the .xlsx price lists change.

Reads every .xlsx in PRICE_LIST_DIR, normalizes + embeds each row's
description (via embedding_service, so it matches the live-query path
exactly), and replaces the entire catalog_items table (TRUNCATE + insert).
Simplest correct option at ~2,500 rows — no diffing/staleness tracking needed.

Usage (from inside rfq_service/, same as dry_run_email.py):
    python embed_price_lists.py
    python embed_price_lists.py --dir /path/to/other/xlsx/folder

Requires: DATABASE_URL, OPENAI_API_KEY (see .env.example). schema.sql must
already be applied to the target database.
"""

import argparse
import glob
import logging
import os

from dotenv import load_dotenv
load_dotenv()

import embedding_service
import db
from text_utils import norm_exact

logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"))
logger = logging.getLogger(__name__)

PRICE_DIR_DEFAULT = os.getenv("PRICE_LIST_DIR", "PRICE_LIST_excel")

# Header keyword -> canonical field. First match wins; order matters.
# (Same mapping price_lookup_service used to own — this script is now the
# only place that reads the raw Excel files.)
_PART_KEYS = ("al item code", "item code", "anytime code", "part", "code")
_DESC_KEYS = ("description", "desc")
_PRICE_KEYS = ("lpl", "list price", "price", "rate")
_LEAD_KEYS = ("lead time", "delivery", "lead")
_HSN_KEYS = ("hsn",)


def _to_price(v):
    if v is None:
        return None
    try:
        return round(float(str(v).replace(",", "").strip()), 2)
    except (ValueError, TypeError):
        return None


def _match_header(cell: str, keys) -> bool:
    c = norm_exact(cell)
    if keys is _PRICE_KEYS and ("dlp" in c or "discount" in c):
        return False
    return any(k in c for k in keys)


def _best_column(cells, keys):
    best_idx, best_rank = None, len(keys)
    for idx, cell in enumerate(cells):
        c = norm_exact(cell)
        if not c:
            continue
        if keys is _PRICE_KEYS and ("dlp" in c or "discount" in c):
            continue
        for rank, k in enumerate(keys):
            if k in c and rank < best_rank:
                best_idx, best_rank = idx, rank
                break
    return best_idx


def _dedupe_colmap(colmap, cells):
    seen = {}
    for field in list(colmap.keys()):
        idx = colmap[field]
        if idx in seen:
            other = seen[idx]
            if len(norm_exact(cells[idx])) and field == "part_number":
                del colmap[other]
                seen[idx] = field
            else:
                del colmap[field]
        else:
            seen[idx] = field


def _find_header(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 8:
            break
        cells = [("" if c is None else str(c)) for c in row]
        if not any(cells):
            continue
        colmap = {}
        for field, keys in (
            ("part_number", _PART_KEYS), ("description", _DESC_KEYS),
            ("price", _PRICE_KEYS), ("lead_time", _LEAD_KEYS), ("hsn", _HSN_KEYS),
        ):
            idx = _best_column(cells, keys)
            if idx is not None:
                colmap[field] = idx
        _dedupe_colmap(colmap, cells)
        if ("description" in colmap or "part_number" in colmap) and "price" in colmap:
            return colmap, i
    return {}, -1


def _row_to_record(row, colmap, fname, sheet):
    def cell(field):
        idx = colmap.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    part = cell("part_number")
    desc = cell("description")
    if (part is None or str(part).strip() == "") and (desc is None or str(desc).strip() == ""):
        return None  # trailing empty row (File 1 pads to ~1M rows)

    desc_str = str(desc).strip() if desc is not None else None
    part_str = str(part).strip() if part is not None else None
    return {
        "source_file": fname,
        "sheet": sheet,
        "part_number": part_str,
        "norm_part": norm_exact(part_str),
        "description": desc_str,
        "norm_description": norm_exact(desc_str),
        "size_token": embedding_service.extract_size_token(desc_str or ""),
        "price": _to_price(cell("price")),
        "lead_time": (str(cell("lead_time")).strip() if cell("lead_time") is not None else None),
        "hsn": (str(cell("hsn")).strip() if cell("hsn") is not None else None),
        "_embed_text": embedding_service.normalize_for_embedding(desc_str or ""),
    }


def read_excel_rows(directory: str) -> list[dict]:
    import openpyxl

    rows: list[dict] = []
    files = sorted(glob.glob(os.path.join(directory, "*.xlsx")))
    if not files:
        logger.warning("[INGEST] no .xlsx found in %s", directory)

    for path in files:
        fname = os.path.basename(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        file_rows = 0
        for ws in wb.worksheets:
            colmap, header_row_idx = _find_header(ws)
            if not colmap or ("description" not in colmap and "part_number" not in colmap):
                continue
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i <= header_row_idx:
                    continue
                rec = _row_to_record(row, colmap, fname, ws.title)
                if rec is None:
                    continue
                rows.append(rec)
                file_rows += 1
        wb.close()
        logger.info("[INGEST] parsed %d rows from %s", file_rows, fname)

    return rows


def ingest(directory: str) -> int:
    rows = read_excel_rows(directory)
    if not rows:
        logger.warning("[INGEST] nothing to ingest — aborting without touching the table")
        return 0

    logger.info("[INGEST] embedding %d descriptions...", len(rows))
    vectors = embedding_service.embed_batch([r["_embed_text"] for r in rows])

    conn = db.get_conn()
    with conn.cursor() as cur:
        cur.execute("TRUNCATE catalog_items")
        cur.executemany(
            "INSERT INTO catalog_items (source_file, sheet, part_number, norm_part, "
            "description, norm_description, size_token, price, lead_time, hsn, embedding) "
            "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
            [
                (rec["source_file"], rec["sheet"], rec["part_number"], rec["norm_part"],
                 rec["description"], rec["norm_description"], rec["size_token"],
                 rec["price"], rec["lead_time"], rec["hsn"], vec)
                for rec, vec in zip(rows, vectors)
            ],
        )
        cur.execute("ANALYZE catalog_items")
    logger.info("[INGEST] loaded %d rows into catalog_items", len(rows))
    return len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=PRICE_DIR_DEFAULT, help="directory of .xlsx price lists")
    args = ap.parse_args()
    n = ingest(args.dir)
    print(f"Ingested {n} catalog rows from {args.dir}")


if __name__ == "__main__":
    main()
